"""
Build a real, downloadable .pptx file from a SlideCraft deck (panel-grid JSON),
drawing native PowerPoint shapes/tables/text so it opens correctly in Canva,
PowerPoint, or Google Slides — this is the actual deliverable file, independent
of Canva's Brand Template API (which this account can't use).
"""
import io
import math
import re

import requests
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.45)
HEADER_H = Inches(0.85)
FOOTER_H = Inches(0.28)
GUTTER = Inches(0.14)


def _hex_to_rgb(hexstr):
    h = (hexstr or "#0f172a").lstrip("#")
    if len(h) != 6:
        h = "0f172a"
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _resolve_color(token, brand):
    if not token:
        return brand["primaryHex"]
    if isinstance(token, str) and token.startswith("#"):
        return token
    return {"primary": brand["primaryHex"], "secondary": brand["secondaryHex"], "accent": brand["accentHex"]}.get(token, brand["primaryHex"])


def _no_line(shape):
    shape.line.fill.background()


def _fill(shape, hexstr):
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex_to_rgb(hexstr)


def _set_text(tf, text, size=11, color="#1e293b", bold=False, align=PP_ALIGN.LEFT, italic=False, font="Calibri"):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run() if not p.runs else p.runs[0]
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = _hex_to_rgb(color)
    return p


def _add_para(tf, text, size=10, color="#334155", bold=False, align=PP_ALIGN.LEFT):
    p = tf.add_paragraph()
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _hex_to_rgb(color)
    return p


def _textbox(slide, x, y, w, h, text, size=11, color="#1e293b", bold=False, align=PP_ALIGN.LEFT, anchor=None, italic=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    _set_text(tf, text, size=size, color=color, bold=bold, align=align, italic=italic)
    return box


def _rect(slide, x, y, w, h, fill_hex=None, line_hex=None, rounded=True, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, x, y, w, h)
    if rounded:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    if fill_hex:
        _fill(shp, fill_hex)
    else:
        shp.fill.background()
    if line_hex:
        shp.line.color.rgb = _hex_to_rgb(line_hex)
        shp.line.width = Pt(0.75)
    else:
        _no_line(shp)
    shp.shadow.inherit = False
    return shp


def _grid_rect(px, py, pw, ph, col, row, colspan, rowspan, cols, rows):
    cw = (pw - GUTTER * (cols - 1)) / cols
    ch = (ph - GUTTER * (rows - 1)) / rows
    x = px + (col - 1) * (cw + GUTTER)
    y = py + (row - 1) * (ch + GUTTER)
    w = cw * colspan + GUTTER * (colspan - 1)
    h = ch * rowspan + GUTTER * (rowspan - 1)
    return x, y, w, h


# ---------------------------------------------------------------------------
# Panel body renderers — each draws directly onto the slide within (x,y,w,h)
# ---------------------------------------------------------------------------

def _body_mini_cards(slide, body, x, y, w, h, brand):
    cards = (body.get("cards") or [])[:4]
    if not cards:
        return
    n = len(cards)
    cols = 2 if n > 1 else 1
    rows = math.ceil(n / cols)
    gap = Inches(0.06)
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for i, card in enumerate(cards):
        cx = x + (i % cols) * (cw + gap)
        cy = y + (i // cols) * (ch + gap)
        _rect(slide, cx, cy, cw, ch, fill_hex="#F8FAFC", line_hex="#E2E8F0")
        tf = slide.shapes.add_textbox(cx + Inches(0.06), cy + Inches(0.03), cw - Inches(0.12), ch - Inches(0.06)).text_frame
        tf.word_wrap = True
        _set_text(tf, card.get("label", ""), size=8, bold=True, color=brand["primaryHex"])
        _add_para(tf, card.get("value", ""), size=8.5, color="#334155")


def _body_bullets(slide, body, x, y, w, h, brand):
    items = (body.get("bullets") or [])[:6]
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    if not items:
        return
    _set_text(tf, f"•  {items[0]}", size=10.5, color="#1e293b")
    for b in items[1:]:
        _add_para(tf, f"•  {b}", size=10.5, color="#1e293b")


def _body_bars(slide, body, x, y, w, h, brand):
    bars = (body.get("bars") or [])[:5]
    if not bars:
        return
    row_h = h / len(bars)
    maxv = max((abs(b.get("value", 0)) for b in bars), default=100) or 100
    for i, b in enumerate(bars):
        ry = y + i * row_h
        _textbox(slide, x, ry, w * 0.4, row_h * 0.5, b.get("label", ""), size=8.5, color="#334155")
        track_x = x
        track_y = ry + row_h * 0.5
        track_w = w
        track_h = min(Inches(0.12), row_h * 0.35)
        _rect(slide, track_x, track_y, track_w, track_h, fill_hex="#E2E8F0", rounded=True)
        val = abs(b.get("value", 0))
        frac = max(0.03, min(1.0, val / maxv))
        _rect(slide, track_x, track_y, track_w * frac, track_h, fill_hex=brand["accentHex"], rounded=True)
        suffix = "%" if b.get("isPercent") else ""
        _textbox(slide, x + w * 0.72, ry, w * 0.28, row_h * 0.5, f"{b.get('value', '')}{suffix}", size=8.5, bold=True, color=brand["primaryHex"], align=PP_ALIGN.RIGHT)


def _body_donut(slide, body, x, y, w, h, brand):
    segs = (body.get("segments") or [])[:5]
    if not segs:
        return
    total = sum(abs(s.get("value", 0)) for s in segs) or 1
    size = min(w, h) * 0.62
    cx = x + Inches(0.05)
    cy = y + (h - size) / 2
    palette = [brand["primaryHex"], brand["accentHex"], brand["secondaryHex"], "#94A3B8", "#CBD5E1"]
    start = 0.0
    for i, s in enumerate(segs):
        frac = abs(s.get("value", 0)) / total
        sweep = frac * 360.0
        pie = slide.shapes.add_shape(MSO_SHAPE.PIE, cx, cy, size, size)
        try:
            pie.adjustments[0] = start
            pie.adjustments[1] = start + sweep
        except Exception:
            pass
        _fill(pie, palette[i % len(palette)])
        _no_line(pie)
        start += sweep
    hole = size * 0.5
    _rect(slide, cx + (size - hole) / 2, cy + (size - hole) / 2, hole, hole, fill_hex="#FFFFFF", rounded=False, line_hex=None)
    legend_x = cx + size + Inches(0.12)
    legend_w = max(Inches(0.4), (x + w) - legend_x)
    ty = cy
    for i, s in enumerate(segs):
        _rect(slide, legend_x, ty + Inches(0.02), Inches(0.09), Inches(0.09), fill_hex=palette[i % len(palette)], rounded=False)
        _textbox(slide, legend_x + Inches(0.14), ty - Inches(0.03), legend_w, Inches(0.2), f"{s.get('label','')} {round(abs(s.get('value',0)))}%", size=7.5, color="#334155")
        ty += min(Inches(0.24), h / max(1, len(segs)))


def _body_table(slide, body, x, y, w, h, brand):
    headers = body.get("headers") or []
    rows = body.get("rows") or []
    if not headers:
        return
    nrows = min(len(rows), 4) + 1
    gshape = slide.shapes.add_table(nrows, len(headers), x, y, w, h)
    table = gshape.table
    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = str(htext)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _hex_to_rgb(brand["primaryHex"])
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
    for ri, row in enumerate(rows[:4]):
        for c, val in enumerate(row[: len(headers)]):
            cell = table.cell(ri + 1, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if ri % 2 == 0 else _hex_to_rgb("#F1F5F9")
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.color.rgb = _hex_to_rgb("#1e293b")


def _body_stat_pair(slide, body, x, y, w, h, brand):
    stats = (body.get("stats") or [])[:3]
    if not stats:
        return
    cw = w / len(stats)
    for i, s in enumerate(stats):
        cx = x + i * cw
        _textbox(slide, cx, y, cw, h * 0.55, str(s.get("value", "")), size=20, bold=True, color=brand["primaryHex"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        _textbox(slide, cx, y + h * 0.55, cw, h * 0.22, s.get("label", ""), size=7.5, bold=True, color="#64748B", align=PP_ALIGN.CENTER)
        if s.get("delta"):
            _textbox(slide, cx, y + h * 0.77, cw, h * 0.2, s.get("delta", ""), size=8.5, bold=True, color=brand["accentHex"], align=PP_ALIGN.CENTER)


def _body_before_after(slide, body, x, y, w, h, brand):
    _textbox(slide, x, y, w, h * 0.25, body.get("label", ""), size=9, bold=True, color="#64748B")
    half = w / 2 - Inches(0.05)
    before_box = _textbox(slide, x, y + h * 0.3, half, h * 0.6, str(body.get("before", "")), size=15, color="#94A3B8", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for p in before_box.text_frame.paragraphs:
        for r in p.runs:
            r.font.strike = "sngStrike"
    _textbox(slide, x + half + Inches(0.1), y + h * 0.3, half, h * 0.6, str(body.get("after", "")), size=17, bold=True, color=brand["primaryHex"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if body.get("delta"):
        up = body.get("direction", "up") != "down"
        _textbox(slide, x, y + h * 0.85, w, Inches(0.2), f"{'▲' if up else '▼'} {body.get('delta')}", size=9, bold=True, color="#16A34A" if up else "#DC2626", align=PP_ALIGN.CENTER)


def _body_flow(slide, body, x, y, w, h, brand):
    steps = (body.get("steps") or [])[:5]
    if not steps:
        return
    vertical = body.get("direction") == "vertical"
    n = len(steps)
    gap = Inches(0.1)
    if vertical:
        bh = (h - gap * (n - 1)) / n
        for i, s in enumerate(steps):
            by = y + i * (bh + gap)
            _rect(slide, x, by, w, bh, fill_hex="#F8FAFC", line_hex=brand["accentHex"])
            tf = slide.shapes.add_textbox(x + Inches(0.05), by, w - Inches(0.1), bh).text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.word_wrap = True
            _set_text(tf, s.get("label", ""), size=8.5, bold=True, color=brand["primaryHex"], align=PP_ALIGN.CENTER)
            if s.get("sublabel"):
                _add_para(tf, s.get("sublabel"), size=7, color="#64748B", align=PP_ALIGN.CENTER)
    else:
        bw = (w - gap * (n - 1)) / n
        for i, s in enumerate(steps):
            bx = x + i * (bw + gap)
            _rect(slide, bx, y, bw, h, fill_hex="#F8FAFC", line_hex=brand["accentHex"])
            tf = slide.shapes.add_textbox(bx + Inches(0.03), y, bw - Inches(0.06), h).text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.word_wrap = True
            _set_text(tf, s.get("label", ""), size=8, bold=True, color=brand["primaryHex"], align=PP_ALIGN.CENTER)
            if s.get("sublabel"):
                _add_para(tf, s.get("sublabel"), size=6.5, color="#64748B", align=PP_ALIGN.CENTER)


def _body_hub_spoke(slide, body, x, y, w, h, brand):
    spokes = (body.get("spokes") or [])[:6]
    cx, cy = x + w / 2, y + h / 2
    hub_w, hub_h = min(w * 0.32, Inches(1.4)), min(h * 0.32, Inches(0.6))
    hub = _rect(slide, cx - hub_w / 2, cy - hub_h / 2, hub_w, hub_h, fill_hex=brand["primaryHex"])
    tf = hub.text_frame
    tf.word_wrap = True
    _set_text(tf, body.get("center", ""), size=8, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)
    n = max(1, len(spokes))
    radius_x, radius_y = w / 2 - Inches(0.55), h / 2 - Inches(0.35)
    node_w, node_h = Inches(1.05), Inches(0.4)
    for i, sp in enumerate(spokes):
        angle = (2 * math.pi * i / n) - math.pi / 2
        nx = cx + radius_x * math.cos(angle)
        ny = cy + radius_y * math.sin(angle)
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(cx), int(cy), int(nx), int(ny))
        conn.line.color.rgb = _hex_to_rgb(brand["accentHex"])
        conn.line.width = Pt(1.25)
        node = _rect(slide, nx - node_w / 2, ny - node_h / 2, node_w, node_h, fill_hex="#F8FAFC", line_hex=brand["accentHex"])
        ntf = node.text_frame
        ntf.word_wrap = True
        _set_text(ntf, sp, size=7.5, bold=True, color=brand["primaryHex"], align=PP_ALIGN.CENTER)


def _body_chevron_phases(slide, body, x, y, w, h, brand):
    phases = (body.get("phases") or [])[:4]
    if not phases:
        return
    n = len(phases)
    gap = Inches(0.03)
    bw = (w - gap * (n - 1)) / n
    for i, ph in enumerate(phases):
        bx = x + i * (bw + gap)
        shp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, bx, y, bw + Inches(0.15), h)
        color = brand["accentHex"] if ph.get("active") else brand["secondaryHex"]
        _fill(shp, color)
        _no_line(shp)
        tf = shp.text_frame
        tf.word_wrap = True
        _set_text(tf, ph.get("label", ""), size=8, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)
        if ph.get("detail"):
            _add_para(tf, ph.get("detail"), size=6.5, color="#FFFFFF", align=PP_ALIGN.CENTER)


ICON_GLYPH = {"heart": "♥", "chart": "▲", "shield": "◆", "users": "◉", "arrows": "↔",
              "clock": "◷", "target": "◎", "star": "★", "globe": "●", "bolt": "⚡"}


def _body_icon_grid(slide, body, x, y, w, h, brand):
    items = (body.get("items") or [])[:6]
    if not items:
        return
    cols = 3 if len(items) > 2 else len(items)
    rows = math.ceil(len(items) / cols)
    gap = Inches(0.06)
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for i, it in enumerate(items):
        cx = x + (i % cols) * (cw + gap)
        cy = y + (i // cols) * (ch + gap)
        _rect(slide, cx, cy, cw, ch, fill_hex="#F8FAFC", line_hex="#E2E8F0")
        tf = slide.shapes.add_textbox(cx, cy + Inches(0.02), cw, ch - Inches(0.04)).text_frame
        tf.word_wrap = True
        glyph = ICON_GLYPH.get(it.get("icon", ""), "•")
        _set_text(tf, glyph, size=13, color=brand["accentHex"], align=PP_ALIGN.CENTER)
        _add_para(tf, it.get("label", ""), size=7, bold=True, color=brand["primaryHex"], align=PP_ALIGN.CENTER)
        if it.get("caption"):
            _add_para(tf, it.get("caption"), size=6, color="#64748B", align=PP_ALIGN.CENTER)


def _body_matrix(slide, body, x, y, w, h, brand):
    quads = (body.get("quadrants") or [])[:4]
    gap = Inches(0.06)
    qw = (w - gap) / 2
    qh = (h - gap) / 2
    palette = [brand["primaryHex"], brand["accentHex"], brand["secondaryHex"], "#64748B"]
    for i in range(min(4, len(quads))):
        q = quads[i]
        qx = x + (i % 2) * (qw + gap)
        qy = y + (i // 2) * (qh + gap)
        _rect(slide, qx, qy, qw, qh, fill_hex="#F8FAFC", line_hex=palette[i % len(palette)])
        tf = slide.shapes.add_textbox(qx + Inches(0.05), qy + Inches(0.02), qw - Inches(0.1), qh - Inches(0.04)).text_frame
        tf.word_wrap = True
        _set_text(tf, q.get("title", ""), size=8, bold=True, color=palette[i % len(palette)])
        for b in (q.get("bullets") or [])[:3]:
            _add_para(tf, f"• {b}", size=6.5, color="#334155")


def _body_timeline(slide, body, x, y, w, h, brand):
    ms = (body.get("milestones") or [])[:6]
    if not ms:
        return
    line_y = y + h * 0.32
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x), int(line_y), int(x + w), int(line_y))
    conn.line.color.rgb = _hex_to_rgb(brand["accentHex"])
    conn.line.width = Pt(1.5)
    conn.line.dash_style = None
    try:
        ln = conn.line._get_or_add_ln()
        d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
        ln.append(d)
    except Exception:
        pass
    n = len(ms)
    cw = w / n
    for i, m in enumerate(ms):
        cx = x + i * cw + cw / 2
        dot = _rect(slide, cx - Inches(0.06), line_y - Inches(0.06), Inches(0.12), Inches(0.12), fill_hex=brand["primaryHex"], rounded=True)
        tf = slide.shapes.add_textbox(x + i * cw, y, cw, h * 0.3).text_frame
        tf.word_wrap = True
        _set_text(tf, m.get("dateLabel", ""), size=7.5, bold=True, color=brand["primaryHex"], align=PP_ALIGN.CENTER)
        tf2 = slide.shapes.add_textbox(x + i * cw, line_y + Inches(0.1), cw, h * 0.5).text_frame
        tf2.word_wrap = True
        _set_text(tf2, m.get("title", ""), size=7, bold=True, color="#1e293b", align=PP_ALIGN.CENTER)
        if m.get("body"):
            _add_para(tf2, m.get("body"), size=6, color="#64748B", align=PP_ALIGN.CENTER)


def _body_photo(slide, body, x, y, w, h, brand):
    url = body.get("imageUrl")
    if url:
        try:
            resp = requests.get(url, timeout=10)
            resp.raise_for_status()
            slide.shapes.add_picture(io.BytesIO(resp.content), x, y, w, h)
            return
        except Exception:
            pass
    _rect(slide, x, y, w, h, fill_hex="#E2E8F0", rounded=False)
    _textbox(slide, x, y, w, h, body.get("imageQuery", "image"), size=8, color="#94A3B8", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _body_statement(slide, body, x, y, w, h, brand):
    tf = slide.shapes.add_textbox(x, y + h * 0.25, w, h * 0.5).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _set_text(tf, body.get("statement", ""), size=26, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)
    if body.get("subtext"):
        _add_para(tf, body.get("subtext"), size=12, color="#FFFFFF", align=PP_ALIGN.CENTER)


BODY_RENDERERS = {
    "mini-cards": _body_mini_cards, "bullets": _body_bullets, "bars": _body_bars,
    "donut": _body_donut, "table": _body_table, "stat-pair": _body_stat_pair,
    "before-after": _body_before_after, "flow": _body_flow, "hub-spoke": _body_hub_spoke,
    "chevron-phases": _body_chevron_phases, "icon-grid": _body_icon_grid, "matrix": _body_matrix,
    "timeline": _body_timeline, "photo": _body_photo, "statement": _body_statement,
}


def _draw_connector_panel(slide, panel, x, y, w, h, brand):
    direction = panel.get("direction", "right")
    dashed = panel.get("style") == "dashed"
    cx, cy = x + w / 2, y + h / 2
    if direction in ("left", "right"):
        x1, y1, x2, y2 = (x, cy, x + w, cy) if direction == "right" else (x + w, cy, x, cy)
    else:
        x1, y1, x2, y2 = (cx, y, cx, y + h) if direction == "down" else (cx, y + h, cx, y)
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    conn.line.color.rgb = _hex_to_rgb(brand["accentHex"])
    conn.line.width = Pt(2)
    conn.line.end_arrowhead = None
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    if dashed:
        d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
        ln.append(d)


def _draw_panel(slide, panel, x, y, w, h, brand):
    if panel.get("kind") == "connector":
        _draw_connector_panel(slide, panel, x, y, w, h, brand)
        return

    body = panel.get("body") or {}
    if panel.get("emphasize"):
        card = _rect(slide, x, y, w, h, fill_hex=brand["primaryHex"])
        renderer = BODY_RENDERERS.get(body.get("type"))
        if renderer:
            renderer(slide, body, x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), h - Inches(0.2), brand)
        return

    card = _rect(slide, x, y, w, h, fill_hex="#FFFFFF", line_hex="#E2E8F0")
    body_y = y
    body_h = h
    if panel.get("headerLabel"):
        header_h = min(Inches(0.32), h * 0.22)
        header = _rect(slide, x, y, w, header_h, fill_hex=_resolve_color(panel.get("headerColor"), brand), rounded=False)
        header.adjustments  # noop
        tf = header.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.06)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _set_text(tf, panel.get("headerLabel", ""), size=7.5, bold=True, color="#FFFFFF")
        body_y = y + header_h + Inches(0.05)
        body_h = h - header_h - Inches(0.05)

    footer = panel.get("footer") or {}
    if footer.get("text"):
        footer_h = min(Inches(0.28), body_h * 0.2)
        body_h -= footer_h + Inches(0.03)
        f_shape = _rect(slide, x + Inches(0.05), body_y + body_h + Inches(0.03), w - Inches(0.1), footer_h, fill_hex="#F1F5F9", rounded=True)
        ftf = f_shape.text_frame
        ftf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _set_text(ftf, footer.get("text", ""), size=7, bold=True, color=_resolve_color(footer.get("color"), brand), align=PP_ALIGN.CENTER)

    renderer = BODY_RENDERERS.get(body.get("type"))
    if renderer:
        try:
            renderer(slide, body, x + Inches(0.08), body_y + Inches(0.03), w - Inches(0.16), body_h - Inches(0.06), brand)
        except Exception:
            pass


def build_pptx(deck: dict) -> bytes:
    brand = deck.get("brand") or {}
    brand = {"primaryHex": "#0f172a", "secondaryHex": "#475569", "accentHex": "#6366f1", **brand}
    slides_data = deck.get("slides") or []

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]

    for sdata in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(255, 255, 255)

        grid = sdata.get("grid") or {"cols": 1, "rows": 1}
        cols = max(1, min(4, grid.get("cols", 1)))
        rows = max(1, min(2, grid.get("rows", 1)))
        panels = sdata.get("panels") or []
        is_statement = len(panels) == 1 and panels[0].get("emphasize")

        content_y = MARGIN
        content_h = SLIDE_H - MARGIN * 2

        if not is_statement:
            if sdata.get("sectionTag"):
                _rect(slide, MARGIN, MARGIN, Inches(1.6), Inches(0.24), fill_hex=brand["secondaryHex"], rounded=True)
                _textbox(slide, MARGIN, MARGIN, Inches(1.6), Inches(0.24), sdata["sectionTag"], size=7, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            title_y = MARGIN + (Inches(0.28) if sdata.get("sectionTag") else 0)
            _textbox(slide, MARGIN, title_y, SLIDE_W - MARGIN * 2, Inches(0.45), sdata.get("title", ""), size=22, bold=True, color=brand["primaryHex"])
            sub_y = title_y + Inches(0.42)
            if sdata.get("subtitle"):
                _textbox(slide, MARGIN, sub_y, SLIDE_W - MARGIN * 2, Inches(0.24), sdata["subtitle"], size=10, italic=True, color="#94A3B8")
                sub_y += Inches(0.22)
            rule_y = sub_y + Inches(0.04)
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(MARGIN), int(rule_y), int(SLIDE_W - MARGIN), int(rule_y))
            line.line.color.rgb = _hex_to_rgb(brand["accentHex"])
            line.line.width = Pt(1)
            content_y = rule_y + Inches(0.12)
            content_h = SLIDE_H - content_y - MARGIN - FOOTER_H - Inches(0.1)

        content_x = MARGIN
        content_w = SLIDE_W - MARGIN * 2

        for panel in panels:
            col = max(1, min(cols, panel.get("col", 1)))
            row = max(1, min(rows, panel.get("row", 1)))
            colspan = max(1, min(cols - col + 1, panel.get("colSpan", 1)))
            rowspan = max(1, min(rows - row + 1, panel.get("rowSpan", 1)))
            x, y, w, h = _grid_rect(content_x, content_y, content_w, content_h, col, row, colspan, rowspan, cols, rows)
            try:
                _draw_panel(slide, panel, x, y, w, h, brand)
            except Exception:
                pass

        if not is_statement:
            _textbox(slide, MARGIN, SLIDE_H - MARGIN - FOOTER_H, content_w * 0.7, FOOTER_H, deck.get("deckTitle", ""), size=7, color="#CBD5E1")

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
